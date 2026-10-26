#!/usr/bin/env python3
"""
AdVantage OS - 4-Slide Major Project Presentation Generator
===========================================================
Generates a polished 16:9 widescreen PowerPoint presentation (.pptx)
covering the 4 required pillars:
1. Requirement Specifications
2. System Architecture
3. Implementation & Engineering Methodology
4. Results, Benchmark Evaluation & Impact
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Palette
    BG_COLOR = RGBColor(11, 15, 25)          # #0B0F19 Dark slate
    CARD_BG = RGBColor(20, 27, 45)           # #141B2D Card navy
    CARD_BORDER = RGBColor(38, 52, 79)       # #26344F Border
    CARD_BG_ACCENT = RGBColor(25, 34, 58)    # Highlight card
    CYAN = RGBColor(56, 189, 248)            # #38BDF8 Accent
    INDIGO = RGBColor(129, 140, 248)         # #818CF8 Accent 2
    EMERALD = RGBColor(52, 211, 153)         # #34D399 Accent Green
    AMBER = RGBColor(251, 191, 36)           # #FBBF24 Accent Yellow
    WHITE = RGBColor(248, 250, 252)          # #F8FAFC Primary text
    SLATE_LIGHT = RGBColor(203, 213, 225)    # #CBD5E1 Body text
    SLATE_MUTED = RGBColor(148, 163, 184)    # #94A3B8 Secondary text
    SLATE_DARK = RGBColor(100, 116, 139)     # #64748B Subdued

    def set_slide_bg(slide):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = BG_COLOR

    def add_header(slide, stage_str, title_str, subtitle_str):
        # Stage Badge
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.35), Inches(2.2), Inches(0.28))
        badge.fill.solid()
        badge.fill.fore_color.rgb = RGBColor(30, 41, 59)
        badge.line.color.rgb = CYAN
        badge.line.width = Pt(1)
        tf_b = badge.text_frame
        tf_b.word_wrap = False
        p_b = tf_b.paragraphs[0]
        p_b.text = stage_str
        p_b.font.size = Pt(9.5)
        p_b.font.bold = True
        p_b.font.color.rgb = CYAN
        p_b.alignment = PP_ALIGN.CENTER

        # Project Tag Right
        tag_box = slide.shapes.add_textbox(Inches(9.5), Inches(0.32), Inches(3.0), Inches(0.3))
        tf_t = tag_box.text_frame
        p_t = tf_t.paragraphs[0]
        p_t.text = "AdVantage OS | Major Project"
        p_t.font.size = Pt(10)
        p_t.font.bold = True
        p_t.font.color.rgb = INDIGO
        p_t.alignment = PP_ALIGN.RIGHT

        # Main Title Box
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.7), Inches(0.75))
        tf = title_box.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title_str
        p1.font.size = Pt(20)
        p1.font.bold = True
        p1.font.color.rgb = WHITE
        
        p2 = tf.add_paragraph()
        p2.text = subtitle_str
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = SLATE_MUTED

    def create_card(slide, left, top, width, height, title="", tag="", border_color=CARD_BORDER, bg_color=CARD_BG):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)
        
        # Inner text frame
        txBox = slide.shapes.add_textbox(left + Inches(0.12), top + Inches(0.1), width - Inches(0.24), height - Inches(0.2))
        tf = txBox.text_frame
        tf.word_wrap = True
        
        if tag:
            p_tag = tf.paragraphs[0]
            p_tag.text = tag.upper()
            p_tag.font.size = Pt(8.5)
            p_tag.font.bold = True
            p_tag.font.color.rgb = CYAN
            p_title = tf.add_paragraph()
        else:
            p_title = tf.paragraphs[0]

        if title:
            p_title.text = title
            p_title.font.size = Pt(12)
            p_title.font.bold = True
            p_title.font.color.rgb = WHITE
        
        return tf

    # =========================================================================
    # SLIDE 1: REQUIREMENT SPECIFICATIONS
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1)
    add_header(s1, "STAGE 01 : REQUIREMENTS", "Requirement Specifications & Domain Problem Statement", 
               "Bridging hyper-local social micro-trends with physical store execution and zero-marginal-cost marketing")

    # Column 1: Problem Statement (Left)
    tf1 = create_card(s1, Inches(0.8), Inches(1.5), Inches(3.7), Inches(5.5), 
                      "1. SME Operational Vulnerabilities", "Domain Problem Statement")
    
    problems = [
        ("The Flatline Trap in Trend Forecasting:", 
         "Standard regressors trained on raw non-stationary levels exhibit >99.5% lag_1 collinearity, causing 30-day recursive forecasts to collapse into flatlines."),
        ("Sparse & Intermittent Retail Demand:", 
         "SME sales feature zero-demand on 60%+ of days. Moving averages fail on bursty count distributions, triggering costly stockouts or dead inventory."),
        ("Disconnected Marketing & Stockouts:", 
         "Promotional campaigns run blind to store stock. Advertising newly trending products without inventory alignment causes immediate stockouts and wasted ad spend."),
        ("Manual Ingestion Friction:", 
         "SMEs rely on paper receipts, thermal slips, and manual ledger entry, creating 7+ day operational data lags without costly RFID/enterprise POS hardware.")
    ]
    for heading, desc in problems:
        p_h = tf1.add_paragraph()
        p_h.text = f"•  {heading}"
        p_h.font.size = Pt(10)
        p_h.font.bold = True
        p_h.font.color.rgb = AMBER
        
        p_d = tf1.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.size = Pt(9)
        p_d.font.color.rgb = SLATE_LIGHT

    # Column 2: Functional Requirements (Middle)
    tf2 = create_card(s1, Inches(4.7), Inches(1.5), Inches(4.3), Inches(5.5), 
                      "2. Core Functional Requirements", "System Functional Scope (FR)")

    f_reqs = [
        ("FR1: Multi-Source Social Ingestion & NLP:", 
         "Concurrently harvest Twitter/X & Reddit; filter topic noise via Cohere LLM; compute dual sentiment (70% DistilBERT + 30% VADER)."),
        ("FR2: Stationary Velocity Trend Forecaster:", 
         "Predict daily velocity differencing (Δy_t) over 30 days with momentum damping (φ=0.97^(t-1)) and expanding 95% confidence intervals."),
        ("FR3: Autonomous Intermittent Demand Engine:", 
         "50/50 blend of Poisson HistGBR and Teunter-Syntetos-Babai (TSB) with stockout masking to calculate Safety Stock (SS) & Reorder Point (ROP)."),
        ("FR4: Edge Computer Vision Receipt Ingestion:", 
         "Camera-based invoice capture; Tesseract OCR + dual SequenceMatcher/token overlap fuzzy SKU reconciliation to update inventory ledgers."),
        ("FR5: Multimodal Generative Creative Studio:", 
         "One-click text-to-image synthesis (NVIDIA SDXL/Flux) + sub-second ad copy & hashtags (Groq LLaMA-3.3-70b) + K-Means influencer discovery."),
        ("FR6: Proactive Closed-Loop Alerting:", 
         "Continuous cross-domain monitor firing automated WhatsApp notifications when social velocity accelerates as stock approaches ROP.")
    ]
    for heading, desc in f_reqs:
        p_h = tf2.add_paragraph()
        p_h.text = f"•  {heading}"
        p_h.font.size = Pt(9.5)
        p_h.font.bold = True
        p_h.font.color.rgb = CYAN
        
        p_d = tf2.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.size = Pt(8.5)
        p_d.font.color.rgb = SLATE_LIGHT

    # Column 3: Non-Functional Requirements & Constraints (Right)
    tf3 = create_card(s1, Inches(9.2), Inches(1.5), Inches(3.3), Inches(5.5), 
                      "3. Non-Functional Specifications", "Constraints & Quality Attributes")

    nf_reqs = [
        ("Sub-Second Inference Speed:", 
         "Marketing ad copy generated in <1.0s via Groq LLaMA-3.3-70b; trend forecasts generated in <200ms."),
        ("Zero-Disk Cloud Persistence:", 
         "Serialized 894 KB model weights stored as BLOBs in Neon PostgreSQL ml_models table; in-memory hydration on cold boot in <350ms."),
        ("High Service Level Target:", 
         "Calculate dynamic safety stock maintaining >=95% customer service level (Z=1.645) with rolling volatility."),
        ("Graceful Multi-Tier Fallbacks:", 
         "Automatic degradation: Pollinations AI for images, synthetic keyword fallback for social feeds, SQLite for local testing."),
        ("Enterprise Modularity & Security:", 
         "Decoupled FastAPI REST microservices, Pydantic v2 validation, CORS origin isolation, and secure environment vaults.")
    ]
    for heading, desc in nf_reqs:
        p_h = tf3.add_paragraph()
        p_h.text = f"•  {heading}"
        p_h.font.size = Pt(9.5)
        p_h.font.bold = True
        p_h.font.color.rgb = EMERALD
        
        p_d = tf3.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.size = Pt(8.5)
        p_d.font.color.rgb = SLATE_LIGHT


    # =========================================================================
    # SLIDE 2: SYSTEM ARCHITECTURE
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2)
    add_header(s2, "STAGE 02 : ARCHITECTURE", "System Architecture & Closed-Loop Operational Topology", 
               "Bi-directional pipeline synchronizing live social listening, velocity forecasting, inventory, GenAI & edge OCR")

    # Left: Architecture Flowchart Diagram Image
    arch_img_path = "extracted_diagrams/fig1_system_architecture_flowchart.png"
    if os.path.exists(arch_img_path):
        # Card container for image
        img_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(7.2), Inches(5.5))
        img_card.fill.solid()
        img_card.fill.fore_color.rgb = CARD_BG
        img_card.line.color.rgb = CYAN
        img_card.line.width = Pt(1.2)
        
        s2.shapes.add_picture(arch_img_path, Inches(0.9), Inches(1.6), Inches(7.0), Inches(4.9))
        
        # Caption below image inside card
        cap_box = s2.shapes.add_textbox(Inches(0.9), Inches(6.55), Inches(7.0), Inches(0.35))
        p_cap = cap_box.text_frame.paragraphs[0]
        p_cap.text = "Figure 1: AdVantage OS Closed-Loop Operational Architecture & Bi-Directional Subsystems"
        p_cap.font.size = Pt(9)
        p_cap.font.color.rgb = SLATE_MUTED
        p_cap.alignment = PP_ALIGN.CENTER
    else:
        tf_arch = create_card(s2, Inches(0.8), Inches(1.5), Inches(7.2), Inches(5.5), "Architecture Diagram", "Pipeline Overview")

    # Right: Architecture Subsystems Breakdown & Feedback Loop
    tf_layers = create_card(s2, Inches(8.2), Inches(1.5), Inches(4.3), Inches(3.1), 
                            "Layered System Decomposition", "Subsystem Taxonomy")
    
    layers = [
        ("Presentation Layer (React 19 + TypeScript + Vite 8.3):", 
         "Glassmorphic Executive Dashboard, Canvas 95% CI trend trajectories, stock ROP bar charts, and sales donut visualizers."),
        ("Application Gateway (FastAPI Async ASGI):", 
         "Lifespan event manager, DB auto-migrations, memory model hydration, and unified REST endpoints (/predict, /api/demand, /api/inventory)."),
        ("Applied AI & Analytics Engines:", 
         "TrendForecastingModel_v2 (Delta Ensemble), 70/30 DistilBERT/VADER sentiment, Poisson HistGBR + TSB demand planner, K-Means creator matcher."),
        ("Generative AI & LLM Cluster:", 
         "NVIDIA API Catalog (SDXL/Flux), Groq LLaMA-3.3-70b (<1s copywriter), Cohere Command-R+, Pollinations AI fallback."),
        ("Cloud Persistence & Sync:", 
         "Neon Serverless PostgreSQL (pooled), zero-disk BLOB model synchronization, local SQLite fallback.")
    ]
    for h, d in layers:
        p = tf_layers.add_paragraph()
        p.text = f"• {h} {d}"
        p.font.size = Pt(8.2)
        p.font.color.rgb = SLATE_LIGHT

    # Right Bottom: The Closed-Loop Feedback Engine
    tf_loop = create_card(s2, Inches(8.2), Inches(4.75), Inches(4.3), Inches(2.25), 
                          "The Autonomous Closed-Loop Engine", "Self-Correcting Arbitrage", 
                          border_color=EMERALD, bg_color=CARD_BG_ACCENT)
    
    steps = [
        ("1. Social Sensing:", "Harvest Twitter/Reddit signals -> filter with Cohere -> compute dual sentiment."),
        ("2. Velocity Forecasting:", "Forecast 30-day adoption trajectory with momentum damping φ=0.97."),
        ("3. Retail Replenishment:", "Align trend velocity with physical stock -> compute Safety Stock & ROP."),
        ("4. Multimodal Generation:", "Synthesize ad visuals & LLaMA-3.3 ad copy for top trending items."),
        ("5. Point-of-Sale Edge OCR:", "Scan physical receipts via mobile camera -> update inventory ledger."),
        ("6. Proactive Stockout Protection:", "Auto-suppress ad campaigns & send WhatsApp alerts if stock drops near ROP.")
    ]
    for s_title, s_desc in steps:
        p = tf_loop.add_paragraph()
        p.text = f"▶ {s_title} {s_desc}"
        p.font.size = Pt(8.0)
        p.font.color.rgb = SLATE_LIGHT


    # =========================================================================
    # SLIDE 3: IMPLEMENTATION & METHODOLOGY
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3)
    add_header(s3, "STAGE 03 : IMPLEMENTATION", "Implementation & Engineering Methodology", 
               "Applied mathematical formulations, ML algorithmic pipelines, and full-stack software integration")

    # Column 1: Mathematical & Algorithmic Formulations (Left)
    tf_math = create_card(s3, Inches(0.8), Inches(1.5), Inches(4.2), Inches(5.5), 
                          "1. Mathematical Formulations", "Core Algorithms & Formulations")
    
    math_items = [
        ("Velocity Differencing Core (Overcoming Flatline):",
         "Δy_t = y_t - y_{t-1}\n"
         "Instead of regressing on raw levels (where lag_1 has 99.5% collinearity), the model predicts rate-of-change velocity across multi-step recursions."),
        ("Momentum Damping & Confidence Bounds:",
         "Δ_{combined, t} = (0.65 · Δ_{ML, t} + 0.35 · slope_{observed}) · φ^{t-1}\n"
         "Empirical damping factor φ = 0.97 simulates natural consumer adoption S-curves. Trajectory bounded by expanding 95% CI: ±1.96 · σ · √(t / 3.0)."),
        ("Blended Poisson-TSB Retail Replenishment:",
         "Intermittent sales modeled via 50/50 blend of HistGBR (Poisson count loss) + Teunter-Syntetos-Babai (TSB) separating demand probability updating (p_t) and demand size updating (z_t)."),
        ("Dynamic Safety Stock & Stockout Masking:",
         "Safety Stock SS = ⌈1.645 · σ_{56} · √L⌉, ROP = ⌈d · L + SS⌉\n"
         "Stockout Masking: Freezes demand estimation when stock=0 to eliminate artificial zero-sales deflation."),
        ("Dual Computer Vision Fuzzy Alignment:",
         "Score = max(Ratio_{seq}, 0.5 · Ratio_{seq} + 0.5 · Overlap_{tokens})\n"
         "Boosted by 0.85 exact substring matching; threshold >= 0.40 reconciles noisy OCR text to product catalog.")
    ]
    for h, d in math_items:
        p_h = tf_math.add_paragraph()
        p_h.text = f"• {h}"
        p_h.font.size = Pt(9.5)
        p_h.font.bold = True
        p_h.font.color.rgb = CYAN
        
        p_d = tf_math.add_paragraph()
        p_d.text = f"   {d}"
        p_d.font.size = Pt(8.2)
        p_d.font.color.rgb = SLATE_LIGHT

    # Column 2: Key Subsystem Implementations (Middle)
    tf_impl = create_card(s3, Inches(5.2), Inches(1.5), Inches(4.1), Inches(5.5), 
                          "2. Subsystem Implementations", "Applied Pipelines & Services")

    impl_items = [
        ("Zero-Disk Cloud Database Model Hydration:",
         "Model weights (894 KB serialized .pkl bundle) are stored directly inside Neon PostgreSQL ml_models table as BLOBs. On server boot, lifespan hook loads weights into RAM in <350ms, enabling zero-disk serverless deployment."),
        ("Dual Ensemble Sentiment Pipeline:",
         "Combines 70% deep contextual transformer polarity (DistilBERT fine-tuned on SST-2) with 30% lexical heuristic rules (VADER) to accurately score informal social slang, punctuation, and emojis."),
        ("Generative Multimodal Studio:",
         "Connects trending keywords to NVIDIA API Catalog (SDXL / Flux.1-dev) for product imagery, and Groq Cloud (LLaMA-3.3-70b-versatile) for <1s high-converting ad copy and hashtags."),
        ("Micro-Influencer Discovery Engine:",
         "Scrapes creator profiles into 4D feature vectors (subscribers, views, engagement rate, keyword affinity). Applies K-Means clustering and rule-based bot fraud filters."),
        ("Edge Computer Vision Receipt Scanner:",
         "Pytesseract 5.x wrapper + PIL preprocessing (grayscale, contrast boost, adaptive thresholding) extracts tabular receipt data with zero manual bookkeeping.")
    ]
    for h, d in impl_items:
        p_h = tf_impl.add_paragraph()
        p_h.text = f"• {h}"
        p_h.font.size = Pt(9.5)
        p_h.font.bold = True
        p_h.font.color.rgb = INDIGO
        
        p_d = tf_impl.add_paragraph()
        p_d.text = f"   {d}"
        p_d.font.size = Pt(8.2)
        p_d.font.color.rgb = SLATE_LIGHT

    # Column 3: Tech Stack & Full-Stack Integration (Right)
    tf_tech = create_card(s3, Inches(9.5), Inches(1.5), Inches(3.0), Inches(5.5), 
                          "3. Software Stack", "Production Tooling")

    tech_categories = [
        ("Backend Services:", 
         "Python 3.10+, FastAPI 0.143, Uvicorn ASGI, Pydantic v2 schemas, SQLAlchemy 2.1 ORM, Alembic migrations."),
        ("Machine Learning & AI:", 
         "Scikit-learn 1.9, PyTorch 2.12, Hugging Face Transformers (DistilBERT), NLTK VADER, SciPy, Pandas 3.0, NumPy 2.4."),
        ("LLM & Generative APIs:", 
         "Groq Cloud (LLaMA-3.3-70b), NVIDIA Cloud Catalog (Flux/SDXL), Cohere Command-R+, Pollinations AI."),
        ("Computer Vision:", 
         "Tesseract OCR 5.x, Pillow 12.2, difflib fuzzy SequenceMatcher."),
        ("Frontend & Visuals:", 
         "React 19.2, TypeScript 5.x, Vite 8.3, HTML5 Canvas API, Lucide Icons, Custom CSS tokens."),
        ("Database Infrastructure:", 
         "Neon PostgreSQL Serverless (pooled), pgvector, zero-disk bytea storage.")
    ]
    for cat, libs in tech_categories:
        p_c = tf_tech.add_paragraph()
        p_c.text = f"• {cat}"
        p_c.font.size = Pt(9)
        p_c.font.bold = True
        p_c.font.color.rgb = EMERALD
        
        p_l = tf_tech.add_paragraph()
        p_l.text = f"   {libs}"
        p_l.font.size = Pt(8.0)
        p_l.font.color.rgb = SLATE_LIGHT


    # =========================================================================
    # SLIDE 4: RESULTS & BENCHMARK EVALUATION
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4)
    add_header(s4, "STAGE 04 : RESULTS & EVALUATION", "Empirical Results, ML Benchmark & Operational Gains", 
               "Validation across 13,825 test samples, sub-second inference latencies, and real-world SME business impact")

    # Left: Benchmark Visualization Image
    bench_img_path = "extracted_diagrams/fig2_trend_model_benchmark.png"
    if os.path.exists(bench_img_path):
        img_card4 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(7.2), Inches(3.4))
        img_card4.fill.solid()
        img_card4.fill.fore_color.rgb = CARD_BG
        img_card4.line.color.rgb = CYAN
        img_card4.line.width = Pt(1.2)
        
        s4.shapes.add_picture(bench_img_path, Inches(0.9), Inches(1.6), Inches(7.0), Inches(2.9))
        
        cap4 = s4.shapes.add_textbox(Inches(0.9), Inches(4.55), Inches(7.0), Inches(0.3))
        p_c4 = cap4.text_frame.paragraphs[0]
        p_c4.text = "Figure 2: Empirical Performance of TrendForecastingModel_v2 (30-day 95% CI & Benchmark Gains)"
        p_c4.font.size = Pt(8.5)
        p_c4.font.color.rgb = SLATE_MUTED
        p_c4.alignment = PP_ALIGN.CENTER

    # Left Bottom: Comparative Metric Table
    table_card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.0), Inches(7.2), Inches(2.0))
    table_card.fill.solid()
    table_card.fill.fore_color.rgb = CARD_BG
    table_card.line.color.rgb = CARD_BORDER
    table_card.line.width = Pt(1)

    t_box = s4.shapes.add_textbox(Inches(0.9), Inches(5.05), Inches(7.0), Inches(0.3))
    p_t = t_box.text_frame.paragraphs[0]
    p_t.text = "HOLD-OUT TEST BENCHMARK (13,825 SAMPLES)"
    p_t.font.size = Pt(9)
    p_t.font.bold = True
    p_t.font.color.rgb = CYAN

    # Add Table
    table_shape = s4.shapes.add_table(3, 4, Inches(0.95), Inches(5.35), Inches(6.9), Inches(1.4))
    table = table_shape.table
    table.columns[0].width = Inches(2.3)
    table.columns[1].width = Inches(1.4)
    table.columns[2].width = Inches(1.7)
    table.columns[3].width = Inches(1.5)

    headers = ["Evaluation Metric", "Naive Baseline", "AdVantage OS Model", "Performance Gain"]
    rows = [
        ["Mean Absolute Error (MAE)", "0.3326", "0.2456", "+26.15% Improvement"],
        ["R² Score (Variance Explained)", "0.9927", "0.9971", "Near-Optimal Fit (0.997)"]
    ]

    for col_idx, text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(30, 41, 59)
        p = cell.text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER

    for row_idx, row_data in enumerate(rows):
        for col_idx, text in enumerate(row_data):
            cell = table.cell(row_idx + 1, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(17, 24, 39)
            p = cell.text_frame.paragraphs[0]
            p.text = text
            p.font.size = Pt(8.5)
            p.font.bold = (col_idx == 2 or col_idx == 3)
            p.font.color.rgb = EMERALD if col_idx == 3 else (CYAN if col_idx == 2 else SLATE_LIGHT)
            p.alignment = PP_ALIGN.CENTER if col_idx > 0 else PP_ALIGN.LEFT

    # Right Top: Computational Latencies Card
    tf_lat = create_card(s4, Inches(8.2), Inches(1.5), Inches(4.3), Inches(2.6), 
                         "Computational Speed & Latency Profiling", "System Throughput & Scalability")

    latencies = [
        ("Marketing Ad Copy Synthesis (Groq LLaMA-3.3-70b):", "< 1.0 second", 
         "Sub-second generation of promotional captions, tone variants, and hashtag bundles."),
        ("Product Visual Synthesis (NVIDIA SDXL / Flux):", "< 4.5 seconds", 
         "High-fidelity photorealistic ad visuals synthesized on-demand with zero local GPU load."),
        ("Camera Receipt OCR & Fuzzy Reconciliation:", "< 2.2 seconds", 
         "Complete end-to-end receipt parsing, invoice regex parsing, and catalog stock update."),
        ("Zero-Disk Serverless DB Model Hydration:", "< 350 ms", 
         "Deserializes 894 KB model weights from Neon PostgreSQL bytea table directly into RAM.")
    ]
    for h, v, d in latencies:
        p_h = tf_lat.add_paragraph()
        p_h.text = f"• {h}  [{v}]"
        p_h.font.size = Pt(8.5)
        p_h.font.bold = True
        p_h.font.color.rgb = WHITE
        
        p_d = tf_lat.add_paragraph()
        p_d.text = f"   {d}"
        p_d.font.size = Pt(7.8)
        p_d.font.color.rgb = SLATE_LIGHT

    # Right Bottom: Real-World SME Business Impact Card
    tf_impact = create_card(s4, Inches(8.2), Inches(4.25), Inches(4.3), Inches(2.75), 
                            "Real-World SME Business Impact", "Commercial Efficacy", 
                            border_color=EMERALD, bg_color=CARD_BG_ACCENT)

    impacts = [
        ("Stockout Risk Reduction:", 
         "Simulations indicate up to an 85% reduction in stockout cascades on viral products due to proactive ROP warnings and lead-time buffering."),
        ("Elimination of Dead Inventory:", 
         "TSB demand decay suppresses over-ordering on fading trends, preventing cash lockup in unsellable stock."),
        ("Marketing Arbitrage & Capital Protection:", 
         "Automatic ad suppression on low-stock SKUs prevents negative-ROI ad spend; creative design agency costs drop from $500+/campaign to near $0."),
        ("Operational Labor Savings:", 
         "Replaces 10+ hours per week of manual inventory ledger entry and trend guesswork with automated OCR and dashboard intelligence.")
    ]
    for h, d in impacts:
        p_h = tf_impact.add_paragraph()
        p_h.text = f"✔ {h}"
        p_h.font.size = Pt(8.5)
        p_h.font.bold = True
        p_h.font.color.rgb = EMERALD
        
        p_d = tf_impact.add_paragraph()
        p_d.text = f"   {d}"
        p_d.font.size = Pt(7.8)
        p_d.font.color.rgb = SLATE_LIGHT

    # Save Presentation
    output_path = "AdVantage_OS_Presentation.pptx"
    prs.save(output_path)
    print(f"[SUCCESS] 4-Slide PowerPoint presentation generated successfully: {output_path}")

if __name__ == "__main__":
    create_deck()
